"""UI Flet en mode web headless pour test visuel + screenshots. Jetable.
État piloté par la variable d'env UISTATE : empty | table | dark | dialog |
partis (partis masqués) | partis_tous (recherche ne trouvant que des partis masqués) |
readme | readme_dialog | readme_dark (captures du README : bailleur + locataires remplis) |
reprise_proposer | reprise_choisir | reprise_premier (question avant de générer) |
rapport | rapport_ecarts (compte rendu après génération)."""
import os
from pathlib import Path

import flet as ft
import interface_flet as gui

STATE = os.environ.get("UISTATE", "empty")
SEED = [
    {"nom": "Martin", "prenom": "Alice", "identifiant": "A101",
     "type_bien": "Appartement", "adresse": "12 rue des Lilas, 75011 Paris",
     "loyer_nu": 800, "charges": 150, "part_caf": 200, "depot_garantie": 1600,
     "date_entree": "2024-01-15", "date_sortie": ""},
    {"nom": "Bernard", "prenom": "Sophie", "identifiant": "M2",
     "type_bien": "Maison", "adresse": "8 avenue du Général Leclerc, 92100 Boulogne",
     "loyer_nu": 1200, "charges": 200, "part_caf": 0, "depot_garantie": 2400,
     "date_entree": "2023-06-01", "date_sortie": "2024-12-31",
     "caution_rendue": True, "observation": "Fin de bail"},
    {"nom": "Nguyen", "prenom": "Thomas", "identifiant": "B3",
     "type_bien": "Studio", "adresse": "45 boulevard Voltaire, 75011 Paris",
     "loyer_nu": 650, "charges": 90, "part_caf": 350, "depot_garantie": 650,
     "date_entree": "2025-03-01", "date_sortie": ""},
]


def main(page):
    app = gui.AppLoyers(page)
    if STATE in ("table", "dark", "dialog", "tri", "search", "partis", "partis_tous",
                 "readme", "readme_dialog", "readme_dark"):
        app.locataires = list(SEED)
        app._rafraichir_table()
    if STATE == "tri":
        app._tri_cle, app._tri_asc = "loyer", False  # loyer décroissant
        app._rafraichir_table()
    if STATE == "search":
        app._recherche = "paris"
        app._rafraichir_table()
    if STATE in ("partis", "partis_tous"):
        app._afficher_partis = app._switch_partis.value = False
        app._recherche = "sophie" if STATE == "partis_tous" else ""
        app._rafraichir_table()
    if STATE.startswith("readme"):
        app.b_nom.value, app.b_sci.value = "Les Tilleuls", True
        app.b_sci_nom.value, app.b_sci_nom.disabled = "SCI Les Tilleuls", False
        app.b_adresse.value = "12 rue des Lilas, 35000 Rennes"
        app.b_tel.value, app.b_email.value = "06 12 34 56 78", "contact@les-tilleuls.fr"
        page.update()
    if STATE == "readme_dark":
        app.apparence = "sombre"
        app._appliquer_apparence()
    if STATE == "readme_dialog":
        import threading
        threading.Timer(1.2, lambda: gui.DialogueLocataire(
            page, app._modules_courants(), app._adresses_connues(), SEED[1],
            app._on_loc_valide(1)).ouvrir()).start()
    if STATE == "dark":
        app.apparence = "sombre"
        app._appliquer_apparence()
    if STATE == "dialog":
        import threading
        threading.Timer(1.2, lambda: gui.DialogueLocataire(
            page, app._modules_courants(), app._adresses_connues(), None,
            app._on_loc_valide(None)).ouvrir()).start()
    if STATE == "reglages":
        import threading
        threading.Timer(3.5, lambda: app.page.show_dialog(app._dlg_reglages)).start()
    # Reprise des saisies : questions posées avant de générer, puis compte rendu.
    ancien = Path("/data/Suivi_Les_Tilleuls AU 28-09-2026.xlsx")
    nouveau = Path("/data/Suivi_Les_Tilleuls MISE A JOUR 29-10-26.xlsx")
    if STATE.startswith("reprise_"):
        import threading
        cas = STATE.removeprefix("reprise_")
        n_sortie = {"proposer": None, "choisir": 12, "premier": None}[cas]
        n_prec = 0 if cas == "premier" else 48
        threading.Timer(3.5, lambda: app._demander_reprise(
            cas, nouveau, n_sortie, ancien, n_prec)).start()
    if STATE == "generation":
        # Parcours réel sans la boîte de dialogue native : un classeur avec des saisies,
        # puis une génération sous un NOUVEAU nom qui les reprend (source=).
        import tempfile
        import threading
        import generer_suivi_loyers as moteur
        from openpyxl import load_workbook
        tmp = Path(tempfile.mkdtemp())
        vieux, neuf = tmp / "Suivi AU 28-09-2026.xlsx", tmp / "Suivi MISE A JOUR.xlsx"
        app.b_nom.value = "Les Tilleuls"
        app.locataires = list(SEED)
        cfg = moteur.valider_config(app._config())
        moteur.generer_workbook(cfg, vieux)
        wb = load_workbook(vieux)
        ws = wb["A101 - Martin"]
        entetes = {c.value: c.column for c in ws[4]}    # ligne d'en-tête du suivi
        ws.cell(5, entetes["Part locataire reçue"]).value = 750   # 1er mois d'Alice
        wb.save(vieux)
        app.locataires.append({"nom": "Petit", "prenom": "Léa", "identifiant": "C4",
                               "loyer_nu": 500, "charges": 50, "part_caf": 0,
                               "date_entree": "2026-04-01"})
        app._rafraichir_table()
        app._cfg_valide = moteur.valider_config(app._config())

        def lancer():
            app._lancer_generation(neuf, vieux)
            threading.Timer(4, lambda: print(
                "MEMO", gui._classeur_memorise(app._reglages, app._slug(app._cfg_valide)),
                "| JSON", neuf.with_suffix(".json").is_file(),
                "| VIEUX INTACT", not list(tmp.glob("*.bak-*")),
                "| MOIS", moteur.compter_mois_saisis(neuf), flush=True)).start()
        threading.Timer(3.5, lancer).start()
    if STATE in ("rapport", "rapport_ecarts"):
        import threading
        import generer_suivi_loyers as moteur
        bak = ancien.with_name(ancien.stem + ".bak-20261029-182951.xlsx")
        rapport = moteur.RapportGeneration(source=ancien, sauvegarde=bak, mois_repris=48)
        if STATE == "rapport_ecarts":
            rapport.ecartees = [moteur.Ecartee(
                "Dupont Jean", "locataire absent de la liste (renommé ou supprimé)", 9,
                "mois")]
        threading.Timer(3.5, lambda: app._afficher_rapport(
            ancien, rapport, gui._lignes_rapport(rapport, ancien))).start()


ft.run(main, view=None, host="0.0.0.0", port=8550)
